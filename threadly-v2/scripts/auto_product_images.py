#!/usr/bin/env python3
"""Generate five product views and wire them into Juno's bundled catalogue.

Run from threadly-v2:
  python -m pip install requests pillow
  python scripts/auto_product_images.py --limit 2 --backend mock
  python scripts/auto_product_images.py --limit 20 --backend pollinations
  python scripts/auto_product_images.py --backend openai --limit 10
  python scripts/auto_product_images.py --watch --interval 600 --backend pollinations

Pollinations is the default free hosted backend (internet required; availability/rate
limits are outside this project's control). OpenAI requires OPENAI_API_KEY and may
incur API charges. Local generation requires a CUDA GPU plus torch/diffusers.
Images are written to public/products/generated and public/catalog.json is updated.
"""
from __future__ import annotations
import argparse, base64, io, json, os, re, subprocess, time
from pathlib import Path
from urllib.parse import quote
import requests
from PIL import Image, ImageDraw, ImageOps

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "public"
CATALOG = PUBLIC / "catalog.json"
OUT = PUBLIC / "products" / "generated"
VIEWS = {
    "front": "front view",
    "side": "side profile view",
    "back": "back view",
    "detail": "close-up detail showing material and texture",
    "lifestyle": "realistic lifestyle photo of one person wearing or using the item",
}

def safe(value):
    return re.sub(r"[^A-Za-z0-9._-]+", "-", str(value or "item")).strip("-._")[:90] or "item"

def prompt_for(product, view):
    name = str(product.get("name") or product.get("Product Name") or "fashion item")
    cat = str(product.get("cat") or product.get("category") or product.get("Category") or "")
    colour = str(product.get("colour") or product.get("color") or "")
    return (f"Photorealistic ecommerce product photography of {name}, {colour}, category {cat}; "
            f"{VIEWS[view]}; exactly one product, centered, clean studio lighting, high detail, "
            "portrait 2:3 composition, no text, no watermark, no collage.").replace("  ", " ")

class Pollinations:
    def __init__(self, delay=3):
        self.delay = delay
    def __call__(self, prompt, seed):
        url = f"https://image.pollinations.ai/prompt/{quote(prompt)}?width=768&height=1152&model=flux&nologo=true&seed={seed}"
        last = "unknown error"
        for attempt in range(5):
            try:
                r = requests.get(url, timeout=180)
                if r.ok and r.headers.get("content-type", "").startswith("image/"):
                    with Image.open(io.BytesIO(r.content)) as im:
                        result = im.convert("RGB")
                    time.sleep(self.delay)
                    return result
                last = f"HTTP {r.status_code}: {r.text[:120]}"
            except Exception as exc:
                last = str(exc)
            time.sleep(5 * (attempt + 1))
        raise RuntimeError(f"Pollinations failed: {last}")

class OpenAIImages:
    def __init__(self):
        self.key = os.environ.get("OPENAI_API_KEY", "")
        if not self.key:
            raise RuntimeError("Set OPENAI_API_KEY to use the OpenAI backend.")
        self.model = os.environ.get("OPENAI_IMAGE_MODEL", "gpt-image-1")
    def __call__(self, prompt, seed):
        body = {"model": self.model, "prompt": prompt, "size": "1024x1536", "quality": "low", "n": 1}
        r = requests.post("https://api.openai.com/v1/images/generations",
                          headers={"Authorization": f"Bearer {self.key}"}, json=body, timeout=300)
        if not r.ok:
            raise RuntimeError(f"OpenAI image API HTTP {r.status_code}: {r.text[:300]}")
        data = r.json()["data"][0]
        if data.get("b64_json"):
            raw = base64.b64decode(data["b64_json"])
        elif data.get("url"):
            response = requests.get(data["url"], timeout=180)
            response.raise_for_status()
            raw = response.content
        else:
            raise RuntimeError("Image API response contained no image data.")
        with Image.open(io.BytesIO(raw)) as im:
            return im.convert("RGB")

class Local:
    def __init__(self):
        import torch
        from diffusers import AutoPipelineForText2Image
        self.torch = torch
        self.pipe = AutoPipelineForText2Image.from_pretrained(
            "stabilityai/sdxl-turbo", torch_dtype=torch.float16, variant="fp16"
        ).to("cuda")
    def __call__(self, prompt, seed):
        generator = self.torch.Generator("cuda").manual_seed(seed)
        return self.pipe(prompt=prompt, height=768, width=512, num_inference_steps=4,
                         guidance_scale=0.0, generator=generator).images[0].convert("RGB")

class Mock:
    def __call__(self, prompt, seed):
        im = Image.new("RGB", (512, 768), (238, 225, 235))
        ImageDraw.Draw(im).text((16, 16), "TEST IMAGE\n" + prompt[:100], fill=(50, 40, 50), spacing=5)
        return im

def update_workbook(path, products):
    if not path.exists():
        return
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        if "Products" not in wb.sheetnames:
            print(f"Workbook has no Products sheet; skipped: {path}")
            return
        ws = wb["Products"]
        headers = [str(c.value or "").strip() for c in ws[1]]
        def ensure(name):
            if name not in headers:
                headers.append(name)
                ws.cell(1, len(headers), name)
            return headers.index(name) + 1
        id_col = next((i+1 for i,h in enumerate(headers) if h.lower() in ("product id","id","sku")), None)
        image_cols = [ensure(n) for n in ("Image","Image 2","Image 3","Image 4","Image 5")]
        status_col = ensure("Image Status")
        if not id_col:
            print("Workbook has no Product ID/ID/SKU column; skipped workbook sync.")
            return
        by_id = {str(p.get("id")): p for p in products}
        for row in range(2, ws.max_row + 1):
            pid = str(ws.cell(row, id_col).value or "").strip()
            if pid in by_id:
                p = by_id[pid]
                paths = p.get("images", [])
                for col, path_value in zip(image_cols, paths):
                    ws.cell(row, col, path_value)
                if len(paths) == 5:
                    ws.cell(row, status_col, "Complete")
        wb.save(path)
        print(f"Updated workbook image columns: {path}")
    except ImportError:
        print("openpyxl is not installed; skipped optional workbook sync.")

def run(args, backend):
    if not CATALOG.exists():
        raise FileNotFoundError(f"Catalogue not found: {CATALOG}")
    data = json.loads(CATALOG.read_text(encoding="utf-8"))
    products = data.get("products", [])
    if not isinstance(products, list) or not products:
        raise ValueError("catalog.json contains no products.")
    chosen = products[:args.limit] if args.limit else products
    made = 0
    for index, product in enumerate(chosen):
        pid = str(product.get("id") or product.get("sku") or f"item-{index+1}")
        folder = OUT / safe(pid)
        folder.mkdir(parents=True, exist_ok=True)
        paths = []
        failed = False
        for view_index, view in enumerate(VIEWS):
            dest = folder / f"{view}.jpg"
            if not dest.exists():
                try:
                    print(f"Generating {pid} — {view}...", flush=True)
                    image = backend(prompt_for(product, view), seed=1000 + index * 7 + view_index)
                    ImageOps.fit(image, (512, 768)).save(dest, "JPEG", quality=88, optimize=True)
                    made += 1
                except Exception as exc:
                    print(f"FAILED {pid} / {view}: {exc}", flush=True)
                    failed = True
                    continue
            paths.append("/" + dest.relative_to(PUBLIC).as_posix())
        if len(paths) == 5:
            product["images"] = paths
            product["img"] = paths[0]
            product["fallbackImg"] = paths[0]
            product["imageStatus"] = "Complete"
        elif failed:
            product["imageStatus"] = f"Incomplete ({len(paths)}/5 views)"
        CATALOG.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"{pid}: {len(paths)}/5 views available", flush=True)
    update_workbook(args.xlsx, products)
    print(f"Finished: {made} new images for {len(chosen)} products.")
    return made

def git_push():
    subprocess.run(["git", "add", "public/products/generated", "public/catalog.json"], cwd=ROOT, check=True)
    subprocess.run(["git", "commit", "-m", "Add generated Juno product images"], cwd=ROOT, check=True)
    subprocess.run(["git", "push", "origin", "main"], cwd=ROOT, check=True)

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--backend", choices=["pollinations", "openai", "local", "mock"], default="pollinations")
    ap.add_argument("--limit", type=int, default=0, help="first N products; 0 means all")
    ap.add_argument("--watch", action="store_true", help="repeat to pick up new catalogue products")
    ap.add_argument("--interval", type=int, default=600, help="seconds between watch passes")
    ap.add_argument("--xlsx", type=Path, default=ROOT / "Threadly_Database.xlsx")
    ap.add_argument("--git-push", action="store_true", help="commit generated images and catalogue, then push main")
    args = ap.parse_args()
    if args.limit < 0 or args.interval < 1:
        ap.error("--limit must be >= 0 and --interval must be >= 1")
    backend = {"pollinations": Pollinations, "openai": OpenAIImages, "local": Local, "mock": Mock}[args.backend]()
    while True:
        made = run(args, backend)
        if args.git_push and made:
            git_push()
        if not args.watch:
            break
        time.sleep(args.interval)

if __name__ == "__main__":
    main()
